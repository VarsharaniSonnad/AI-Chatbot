import time
import pandas as pd
from bs4 import BeautifulSoup
from docx import Document
from docx.shared import Pt
from selenium import webdriver
from selenium.webdriver.chrome.service import Service
from webdriver_manager.chrome import ChromeDriverManager

# Read all URLs
links = pd.read_csv("website_links.csv")

# Create Word document
document = Document()
document.add_heading("Alphaleo Website Documentation", level=1)

# Open Chrome
options = webdriver.ChromeOptions()

# Comment this if you want Chrome to open visibly
# options.add_argument("--headless")

options.add_argument("--start-maximized")

driver = webdriver.Chrome(
    service=Service(ChromeDriverManager().install()),
    options=options
)

for url in links["URL"]:

    print("=" * 80)
    print("Scraping:", url)

    driver.get(url)

    # Wait for page to load
    time.sleep(4)

    # Scroll to bottom
    driver.execute_script("window.scrollTo(0, document.body.scrollHeight);")
    time.sleep(3)

    # Scroll once more in case lazy-loaded content appears
    driver.execute_script("window.scrollTo(0, document.body.scrollHeight);")
    time.sleep(3)


    html = driver.page_source

    soup = BeautifulSoup(html, "html.parser")

    # Remove unwanted sections
    for tag in soup([
        "script",
        "style",
        "noscript",
        "svg"
    ]):
        tag.decompose()

    title = soup.title.text.strip() if soup.title else url

    # Page title
    document.add_heading(title, level=1)

    document.add_paragraph(url)

    # Find the main content
    main = (
        soup.find("main")
        or soup.find("article")
        or soup.find("body")
    )

    if main is None:
        continue

    # Traverse in order
    for element in main.descendants:

        if getattr(element, "name", None) is None:
            continue

        text = element.get_text(" ", strip=True)

        if len(text) < 3:
            continue

        if element.name == "h1":
            document.add_heading(text, level=1)

        elif element.name == "h2":
            document.add_heading(text, level=2)

        elif element.name == "h3":
            document.add_heading(text, level=3)

        elif element.name == "h4":
            document.add_heading(text, level=4)

        elif element.name == "h5":
            document.add_heading(text, level=5)

        elif element.name == "h6":
            document.add_heading(text, level=6)

        elif element.name == "p":

            p = document.add_paragraph()
            run = p.add_run(text)
            run.font.size = Pt(11)

        elif element.name == "li":

            document.add_paragraph(
                text,
                style="List Bullet"
            )

    document.add_page_break()

driver.quit()

document.save("Alphaleo_Website_Data1.docx")

print("\nDone.")